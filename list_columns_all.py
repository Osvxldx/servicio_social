import openpyxl
import os

def list_columns(path, sheet_name):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        if sheet_name not in wb.sheetnames:
            print(f"Sheet {sheet_name} not found.")
            return
        sheet = wb[sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        print(f"Columns in '{sheet_name}':")
        for i, val in enumerate(header_row):
             if val is not None:
                print(f"{i}: {val}")
    except Exception as e:
        print(f"Error reading excel: {e}")

if __name__ == "__main__":
    path = 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx'
    list_columns(path, 'BASE DE DATOS')
    list_columns(path, 'PAGOS 2024')
    # Let's also find which sheets look like "PAGOS" sheets
    wb = openpyxl.load_workbook(path, read_only=True)
    print(f"All sheets: {wb.sheetnames}")
