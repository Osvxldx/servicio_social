
import openpyxl

def compare_headers(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        
        for name in ['lista usuarios', 'BASE DE DATOS']:
            if name in wb.sheetnames:
                sheet = wb[name]
                print(f"--- Sheet: {name} ---")
                
                # Check first few rows for something that looks like a header
                for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=4, values_only=True), 1):
                    row_list = list(row)
                    # Filter None
                    cleaned = [str(x).strip() for x in row_list if x is not None]
                    if len(cleaned) > 3: # Assuming a valid header has > 3 columns
                         print(f"Row {i} (Candidate Header): {cleaned}")
                         # If we find "NOMBRE", likely the header
                         if any("NOMBRE" in c.upper() for c in cleaned):
                             print("  -> Likely Header")
    except Exception as e:
        print(f"Error reading excel: {e}")

if __name__ == "__main__":
    compare_headers('RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
