import sqlite3
import openpyxl
import os

def analyze_alignment():
    base_dir = os.getcwd()
    excel_path = os.path.join(base_dir, 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
    db_path = os.path.join(base_dir, 'agua_potable.db')
    
    print("--- DATA ANALYSIS ---")
    
    # Check current DB
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    c.execute("SELECT id, numero_usuario, nombre FROM usuarios ORDER BY id ASC LIMIT 10")
    db_users = c.fetchall()
    print("\n[SQLite DB - First 10 IDs]")
    for u in db_users:
        print(f"ID: {u[0]} | Num: {u[1]} | Name: {u[2]}")
        
    conn.close()
    
    # Check Excel
    print("\n[Excel 'BASE DE DATOS' - First 10 Rows (after header)]")
    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb['BASE DE DATOS']
        row_count = 0
        for row in range(2, 12): # Rows 2 to 11
            num = ws.cell(row=row, column=2).value
            name = ws.cell(row=row, column=3).value
            print(f"Row {row}: Num: {num} | Name: {name}")
    else:
        print("Excel not found.")

if __name__ == "__main__":
    analyze_alignment()
