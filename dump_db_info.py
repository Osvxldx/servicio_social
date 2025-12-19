import openpyxl
import json

wb = openpyxl.load_workbook('RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx', data_only=True)
ws = wb['BASE DE DATOS']
headers = [str(c.value) for c in ws[1]]
row2 = [str(c.value) for c in ws[2]]

info = {
    "headers": headers,
    "row2": row2
}

with open('db_info.json', 'w', encoding='utf-8') as f:
    json.dump(info, f, indent=4, ensure_ascii=False)
