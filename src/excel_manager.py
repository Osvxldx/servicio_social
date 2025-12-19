import openpyxl
from openpyxl.cell.cell import MergedCell
import os
from datetime import datetime

class ExcelManager:
    @staticmethod
    def actualizar_usuario_en_db_excel(datos):
        """Sincroniza datos de usuario con la hoja 'BASE DE DATOS'"""
        base_dir = os.getcwd()
        template_path = os.path.join(base_dir, 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
        
        if not os.path.exists(template_path):
            print(f"No se encontró el archivo Excel en: {template_path}")
            return False

        try:
            wb = openpyxl.load_workbook(template_path)
            if 'BASE DE DATOS' not in wb.sheetnames:
                print("Hoja 'BASE DE DATOS' no encontrada")
                wb.close()
                return False
                
            ws = wb['BASE DE DATOS']
            
            target_row = None
            user_number = str(datos.get('numero_usuario', ''))
            
            # Buscar por No. USUARIO (Columna B / 2)
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=2).value
                if val is not None:
                    # Normalizar comparación: quitar .0 si es float en Excel
                    val_str = str(val).split('.')[0] if isinstance(val, (int, float)) else str(val)
                    if val_str == user_number:
                        target_row = row
                        break
            
            # Si no existe, agregar al final
            if not target_row:
                target_row = ws.max_row + 1
                ws.cell(row=target_row, column=2).value = user_number
            
            # Actualizar campos basados en índices de inspección (1-based)
            # B=2, C=3, D=4, E=5, F=6, G=7, H=8, I=9, J=10, K=11, L=12, M=13, N=14, O=15
            ws.cell(row=target_row, column=3).value = datos.get('nombre', '') # NOMBRE (C)
            ws.cell(row=target_row, column=4).value = datos.get('seccionar', 0) # SECCIONAR (D)
            ws.cell(row=target_row, column=5).value = datos.get('t_pozo', 0) # T. POZO (E)
            ws.cell(row=target_row, column=6).value = datos.get('tomas', 1) # TOMA (F)
            ws.cell(row=target_row, column=7).value = datos.get('conagua', 0) # CONAGUA (G)
            ws.cell(row=target_row, column=8).value = datos.get('drenaje', 0) # DRENAJE (H)
            ws.cell(row=target_row, column=9).value = datos.get('hidrantes', 0) # HIDRANTE (I)
            ws.cell(row=target_row, column=10).value = datos.get('inquilinos', 0) # INQUILINO (J)
            ws.cell(row=target_row, column=11).value = datos.get('vacas', 0) # VACAS (K)
            ws.cell(row=target_row, column=13).value = datos.get('direccion', '') # DIRECCION (M)
            
            wb.save(template_path)
            wb.close()
            return True
        except Exception as e:
            print(f"Error actualizando Excel: {e}")
            return False

    @staticmethod
    def abrir_excel_db():
        """Abre el archivo de Excel de la base de datos"""
        base_dir = os.getcwd()
        template_path = os.path.join(base_dir, 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
        
        if os.path.exists(template_path):
            try:
                os.startfile(template_path)
                return True
            except Exception as e:
                print(f"Error al abrir Excel: {e}")
                return False
        return False

    @staticmethod
    def registrar_pago_en_db_excel(user_number, mes, anio, monto):
        """Registra un pago en la hoja 'BASE DE DATOS' en la columna del mes correspondiente"""
        base_dir = os.getcwd()
        template_path = os.path.join(base_dir, 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
        
        if not os.path.exists(template_path):
            return False

        try:
            wb = openpyxl.load_workbook(template_path)
            if 'BASE DE DATOS' not in wb.sheetnames:
                wb.close()
                return False
                
            ws = wb['BASE DE DATOS']
            
            target_row = None
            user_number_str = str(user_number)
            
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=2).value
                if val is not None:
                    # Normalizar comparación
                    val_str = str(val).split('.')[0] if isinstance(val, (int, float)) else str(val)
                    if val_str == user_number_str:
                        target_row = row
                        break
            
            if not target_row:
                wb.close()
                return False
                
            # Columnas de meses: ENERO es O (15), FEBRERO es P (16)...
            # mes (1-12) -> 14 + mes
            col_index = 14 + mes
            
            ws.cell(row=target_row, column=col_index).value = monto
            # Actualizar AÑO (Col N / 14)
            ws.cell(row=target_row, column=14).value = anio
            
            wb.save(template_path)
            wb.close()
            return True
        except Exception as e:
            print(f"Error registrando pago en Excel: {e}")
            return False

    @staticmethod
    def _get_target_cell(ws, coord):
        """Devuelve la celda real (Top-Left si es Merged)"""
        cell = ws[coord]
        if isinstance(cell, MergedCell):
            for merged_range in ws.merged_cells.ranges:
                if coord in merged_range:
                    return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
        return cell

    @staticmethod
    def generar_recibo(datos):
        base_dir = os.getcwd()
        template_path = os.path.join(base_dir, 'RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
        output_path = os.path.join(base_dir, 'recibo_temp.xlsx')

        if not os.path.exists(template_path):
            raise FileNotFoundError(f"No se encontró la plantilla en: {template_path}")

        wb = openpyxl.load_workbook(template_path)
        
        # Intentar con diferentes nombres de hoja de recibo
        sheet_candidates = ['RECIBO 2025', 'RECIBO NUEVO', 'RECIBO']
        ws = None
        for name in sheet_candidates:
            if name in wb.sheetnames:
                ws = wb[name]
                break
        
        if ws is None:
            ws = wb.active

        def write_safe(coord, value):
            target = ExcelManager._get_target_cell(ws, coord)
            target.value = value
            return target

        folio = datos.get('folio', '')
        fecha = datos.get('fecha', '')
        id_usuario = datos.get('id_usuario', '')
        nombre = datos.get('nombre', '')
        direccion = datos.get('direccion', '')
        tomas = datos.get('tomas', '')
        total = datos.get('total', 0.0)
        conceptos = datos.get('conceptos', [])

        write_safe('L2', folio)
        write_safe('L26', folio)
        write_safe('I5', fecha)
        write_safe('I29', fecha)
        write_safe('C6', id_usuario)
        write_safe('C30', id_usuario)
        write_safe('B8', nombre)
        write_safe('B32', nombre)
        write_safe('H8', direccion)
        write_safe('H32', direccion)
        write_safe('B10', tomas)
        write_safe('B34', tomas)
        write_safe('K18', total)
        write_safe('K42', total)

        start_row_original = 12
        start_row_copy = 36
        written_desc_coords_orig = set()
        
        for i, concepto in enumerate(conceptos):
            desc = concepto.get('descripcion', '')
            precio = concepto.get('importe', 0.0)
            
            coord_desc = f'B{start_row_original + i}'
            coord_precio = f'D{start_row_original + i}'
            
            target_desc = ExcelManager._get_target_cell(ws, coord_desc)
            target_precio = ExcelManager._get_target_cell(ws, coord_precio)
            
            if target_desc.coordinate in written_desc_coords_orig:
                current_text = str(target_desc.value) if target_desc.value else ""
                target_desc.value = current_text + "\n" + desc
            else:
                target_desc.value = desc
                written_desc_coords_orig.add(target_desc.coordinate)
            
            target_precio.value = precio

            coord_desc_copy = f'B{start_row_copy + i}'
            coord_precio_copy = f'D{start_row_copy + i}'
            write_safe(coord_desc_copy, desc)
            write_safe(coord_precio_copy, precio)

        try:
            wb.save(output_path)
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = os.path.join(base_dir, f'recibo_temp_{timestamp}.xlsx')
            wb.save(output_path)

        wb.close()
        return os.path.abspath(output_path)
