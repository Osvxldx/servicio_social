
import openpyxl

def inspect_db_sheet(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        name = 'BASE DE DATOS'
        if name in wb.sheetnames:
            sheet = wb[name]
            print(f"--- Sheet: {name} ---")
            for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=5, values_only=True), 1):
                strings = [str(c) for c in row if c is not None]
                print(f"Row {i}: {strings}")
        else:
            print(f"Sheet {name} not found")
    except Exception as e:
        print(f"Error reading excel: {e}")

if __name__ == "__main__":
    inspect_db_sheet('RECIBOS DE AGUA POTABLE NUEVOS2025.xlsx')
